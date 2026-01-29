import os
import uuid
import threading
import json
import base64
import cv2
import numpy as np
import ssl
import socket
import ssl
from pathlib import Path
from datetime import datetime
from pathlib import Path
from werkzeug.utils import secure_filename
from datetime import datetime, timedelta
from flask import Flask, render_template, request, redirect, url_for, jsonify
import ai_model

from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas
from io import BytesIO
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)
app.secret_key = 'secret-key-12345'

# Конфигурации
BASE_DIR = Path(__file__).parent.absolute()
UPLOAD_FOLDER = BASE_DIR / 'static' / 'videos'
DETECTED_FOLDER = BASE_DIR / 'static' / 'detected'
LOST_ITEMS_FOLDER = DETECTED_FOLDER / 'lost_items'
CAMERA_FRAMES_FOLDER = DETECTED_FOLDER / 'camera_frames'
METADATA_FILE = DETECTED_FOLDER / 'items_metadata.json'

# Создаем папки
for folder in [UPLOAD_FOLDER, DETECTED_FOLDER, LOST_ITEMS_FOLDER]:
    folder.mkdir(parents=True, exist_ok=True)

app.config['UPLOAD_FOLDER'] = str(UPLOAD_FOLDER)
app.config['LOST_ITEMS_FOLDER'] = str(LOST_ITEMS_FOLDER)
app.config['ALLOWED_EXTENSIONS'] = {'mp4', 'avi', 'mov', 'wmv', 'flv', 'mkv', 'webm', 'm4v'}
app.config['MAX_CONTENT_LENGTH'] = 500 * 1024 * 1024  # 500MB

# Словари для хранения статусов
processing_status = {}
video_items = {}  # video_id -> [список найденных предметов]
camera_sessions = {}  # session_id -> CameraSession
camera_detectors = {}  # session_id -> AdvancedLostItemDetector
session_last_activity = {}  # session_id -> last_activity_time
items_metadata = {}

# Создаем/загружаем метаданные
def load_metadata():
    """Загрузка метаданных из JSON файла"""
    global items_metadata
    if METADATA_FILE.exists():
        try:
            with open(METADATA_FILE, 'r', encoding='utf-8') as f:
                items_metadata = json.load(f)
        except:
            items_metadata = {}
    else:
        items_metadata = {}

def save_metadata():
    """Сохранение метаданных в JSON файл"""
    with open(METADATA_FILE, 'w', encoding='utf-8') as f:
        json.dump(items_metadata, f, ensure_ascii=False, indent=2)

# Загружаем при старте
load_metadata()

def add_item_metadata(item_info):
    """Добавление метаданных предмета"""
    global items_metadata
    if 'filename' in item_info:
        items_metadata[item_info['filename']] = {
            'class_name': item_info.get('class_name', 'Сумка/Рюкзак'),
            'confidence': item_info.get('confidence', 0.0),
            'timestamp': item_info.get('timestamp', ''),
            'source': 'camera' if item_info.get('video_name') == 'Камера' else 'video',
            'item_id': item_info.get('item_id', '')
        }
        save_metadata()

def remove_item_metadata(filename):
    """Удаление метаданных предмета"""
    global items_metadata
    if filename in items_metadata:
        del items_metadata[filename]
        save_metadata()

def allowed_file(filename):
    return '.' in filename and \
           filename.rsplit('.', 1)[1].lower() in app.config['ALLOWED_EXTENSIONS']

def save_base64_image(base64_string, folder, filename):
    """Сохраняет base64 изображение в файл"""
    try:
        # Убираем префикс "data:image/jpeg;base64,"
        if ',' in base64_string:
            base64_string = base64_string.split(',')[1]
        
        img_data = base64.b64decode(base64_string)
        filepath = Path(folder) / filename
        
        with open(filepath, 'wb') as f:
            f.write(img_data)
        
        return str(filepath)
    except Exception as e:
        print(f"Error saving base64 image: {e}")
        return None

def process_video_async(video_path, video_id, video_name):
    """Асинхронная обработка видео"""
    try:
        processing_status[video_id] = {'status': 'processing', 'progress': 0}
        video_items[video_id] = []  # Инициализируем пустой список для этого видео
        
        def on_lost_item_found(item_info):
            """Callback при обнаружении потерянного предмета"""
            if video_id in video_items:
                video_items[video_id].append(item_info)
                # Сохраняем метаданные
                add_item_metadata(item_info)
        
        def on_progress_update(progress):
            """Callback для обновления прогресса"""
            if video_id in processing_status:
                processing_status[video_id]['progress'] = progress
        
        # Запускаем анализ видео
        lost_items = ai_model.analyze_video(
            video_path=video_path,
            output_folder=app.config['LOST_ITEMS_FOLDER'],
            video_id=video_id,
            progress_callback=on_progress_update,
            lost_item_callback=on_lost_item_found
        )
        
        processing_status[video_id] = {
            'status': 'completed',
            'progress': 100,
            'lost_items': lost_items,
            'video_name': video_name
        }
        
    except Exception as e:
        processing_status[video_id] = {
            'status': 'error',
            'progress': 0,
            'error': str(e)
        }
        print(f"Error processing video {video_id}: {e}")

@app.route('/')
def index():
    """Главная страница"""
    videos = []
    for file in UPLOAD_FOLDER.iterdir():
        if file.is_file() and allowed_file(file.name):
            videos.append({
                'name': file.name,
                'url': f'/static/videos/{file.name}'
            })
    
    # Получаем найденные предметы из метаданных
    lost_items = []
    
    # Загружаем актуальные метаданные
    load_metadata()
    
    for filename, metadata in items_metadata.items():
        # Проверяем, что файл существует
        item_path = LOST_ITEMS_FOLDER / filename
        if item_path.exists():
            lost_items.append({
                'image_url': f'/static/detected/lost_items/{filename}',
                'video_name': 'Камера' if metadata.get('source') == 'camera' else 'Видео',
                'timestamp': metadata.get('timestamp', 'Неизвестно'),
                'item_id': metadata.get('item_id', ''),
                'filename': filename,
                'class_name': metadata.get('class_name', 'Сумка/Рюкзак'),
                'confidence': metadata.get('confidence', 0.0)
            })
        else:
            # Файл удален, удаляем и метаданные
            remove_item_metadata(filename)
    
    # Сортируем по времени (новые сверху)
    lost_items.sort(key=lambda x: x.get('timestamp', ''), reverse=True)
    
    return render_template('index.html', videos=videos, lost_items=lost_items)

@app.route('/upload', methods=['POST'])
def upload_video():
    """Загрузка и анализ видео"""
    if 'video' not in request.files:
        return jsonify({'success': False, 'error': 'No file uploaded'})
    
    file = request.files['video']
    
    if file.filename == '':
        return jsonify({'success': False, 'error': 'No file selected'})
    
    if file and allowed_file(file.filename):
        # Сохраняем видео
        filename = secure_filename(file.filename)
        video_id = str(uuid.uuid4())[:8]
        save_filename = f"{video_id}_{filename}"
        video_path = UPLOAD_FOLDER / save_filename
        
        # Если файл уже существует, добавляем номер
        counter = 1
        while video_path.exists():
            name, ext = os.path.splitext(filename)
            save_filename = f"{video_id}_{name}_{counter}{ext}"
            video_path = UPLOAD_FOLDER / save_filename
            counter += 1
        
        file.save(str(video_path))
        
        # Запускаем анализ в отдельном потоке
        thread = threading.Thread(
            target=process_video_async,
            args=(str(video_path), video_id, filename)
        )
        thread.daemon = True
        thread.start()
        
        return jsonify({
            'success': True,
            'video_id': video_id,
            'video_name': filename,
            'message': 'Видео загружено и отправлено на анализ'
        })
    
    return jsonify({'success': False, 'error': 'Invalid file format'})

@app.route('/analyze_camera_frame', methods=['POST'])
def analyze_camera_frame():
    """Анализ кадра с камеры с дедупликацией"""
    try:
        data = request.json
        if not data or 'image' not in data:
            return jsonify({'success': False, 'error': 'No image data'})
        
        image_data = data['image']
        session_id = data.get('session_id')
        
        if not session_id or session_id not in camera_detectors:
            return jsonify({'success': False, 'error': 'Invalid session'})
        
        # Обновляем время активности
        session_last_activity[session_id] = datetime.now()
        
        # Получаем детектор для этой сессии
        detector = camera_detectors[session_id]
        
        # Конвертируем base64 в numpy array
        if ',' in image_data:
            image_data = image_data.split(',')[1]
        
        img_bytes = base64.b64decode(image_data)
        nparr = np.frombuffer(img_bytes, np.uint8)
        frame = cv2.imdecode(nparr, cv2.IMREAD_COLOR)
        
        if frame is None:
            return jsonify({'success': False, 'error': 'Failed to decode image'})
        
        # Увеличиваем счетчик кадров
        camera_sessions[session_id]['frame_count'] += 1
        
        # Анализируем кадр
        lost_items, annotated_frame = detector.analyze_frame(
            frame=frame,
            session_id=session_id,
            timestamp=None,  # Время будет сгенерировано внутри
            output_folder=app.config['LOST_ITEMS_FOLDER']
        )
        
        # Сохраняем метаданные для каждого найденного предмета
        for item in lost_items:
            add_item_metadata(item)
        
        result = {
            'success': True,
            'lost_items': lost_items,
            'lost_items_count': len(lost_items),
            'session_id': session_id
        }
        
        return jsonify(result)
        
    except Exception as e:
        print(f"Error analyzing camera frame: {e}")
        return jsonify({'success': False, 'error': str(e)})

@app.route('/start_camera_session', methods=['POST'])
def start_camera_session():
    """Начать новую сессию камеры"""
    session_id = str(uuid.uuid4())[:8]
    
    # Создаем детектор для этой сессии
    detector = ai_model.AdvancedLostItemDetector(model_size='medium', device='cuda')
    camera_detectors[session_id] = detector
    
    # Записываем время начала
    session_last_activity[session_id] = datetime.now()
    camera_sessions[session_id] = {
        'started_at': datetime.now().isoformat(),
        'frame_count': 0,
        'lost_items': [],
        'unique_lost_items': set()  # Для отслеживания уникальных предметов
    }
    
    return jsonify({
        'success': True,
        'session_id': session_id,
        'message': 'Сессия камеры начата'
    })

@app.route('/status/<video_id>')
def get_status(video_id):
    """Получение статуса обработки видео"""
    if video_id in processing_status:
        return jsonify(processing_status[video_id])
    return jsonify({'status': 'not_found'}), 404

@app.route('/items/<video_id>')
def get_items(video_id):
    """Получение найденных предметов для видео"""
    if video_id in video_items:
        return jsonify({
            'success': True,
            'items': video_items[video_id]
        })
    return jsonify({'success': False, 'items': []})

@app.route('/delete_video/<filename>', methods=['POST'])
def delete_video(filename):
    """Удаление видео"""
    try:
        safe_filename = secure_filename(filename)
        video_path = UPLOAD_FOLDER / safe_filename
        if video_path.exists():
            video_path.unlink()
            
            # Удаляем связанные обнаруженные предметы
            video_prefix = safe_filename.split('.')[0]
            for item in LOST_ITEMS_FOLDER.iterdir():
                if item.stem.startswith(video_prefix):
                    item.unlink()
    except Exception as e:
        print(f"Error deleting video: {e}")
    
    return redirect('/')

@app.route('/delete_lost_item/<filename>', methods=['POST'])
def delete_lost_item(filename):
    """Удаление изображения потерянного предмета"""
    try:
        safe_filename = secure_filename(filename)
        
        # Удаляем из LOST_ITEMS_FOLDER
        item_path = LOST_ITEMS_FOLDER / safe_filename
        if item_path.exists():
            item_path.unlink()
            
        # Удаляем метаданные
        remove_item_metadata(safe_filename)
            
    except Exception as e:
        print(f"Error deleting lost item: {e}")
    
    return redirect('/')

@app.route('/cleanup_old_sessions', methods=['POST'])
def cleanup_old_sessions():
    """Очистка старых сессий (вызывается периодически)"""
    now = datetime.now()
    sessions_to_remove = []
    
    for session_id, last_activity in session_last_activity.items():
        if (now - last_activity) > timedelta(hours=1):  # 1 час неактивности
            sessions_to_remove.append(session_id)
    
    for session_id in sessions_to_remove:
        if session_id in camera_detectors:
            del camera_detectors[session_id]
        if session_id in session_last_activity:
            del session_last_activity[session_id]
        if session_id in camera_sessions:
            del camera_sessions[session_id]
    
    return jsonify({
        'success': True,
        'cleaned_sessions': len(sessions_to_remove)
    })

@app.route('/api/save_item_metadata', methods=['POST'])
def save_item_metadata():
    """Сохранение метаданных найденного предмета"""
    try:
        data = request.json
        add_item_metadata(data)
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)})

def get_ip_address():
    """Получить локальный IP адрес"""
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except:
        return "127.0.0.1"

@app.route('/generate_report')
def generate_report():
    """Генерация Excel отчета со статистикой"""
    try:
        # Загружаем метаданные
        load_metadata()
        
        # Создаем новую книгу Excel
        wb = Workbook()
        ws = wb.active
        ws.title = "Статистика"
        
        # Стили
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill(start_color="3498db", end_color="3498db", fill_type="solid")
        title_font = Font(bold=True, size=14)
        center_alignment = Alignment(horizontal="center", vertical="center")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # 1. ЗАГОЛОВОК
        ws.merge_cells('A1:E1')
        ws['A1'] = "ОТЧЕТ ПО ПОТЕРЯННЫМ ВЕЩАМ"
        ws['A1'].font = Font(bold=True, size=16)
        ws['A1'].alignment = center_alignment
        
        ws.merge_cells('A2:E2')
        ws['A2'] = "Система обнаружения потерянных вещей"
        ws['A2'].font = Font(size=12, color="666666")
        ws['A2'].alignment = center_alignment
        
        ws.merge_cells('A3:E3')
        ws['A3'] = f"Отчет сгенерирован: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws['A3'].font = Font(size=10, color="999999")
        ws['A3'].alignment = center_alignment
        
        # 2. ОБЩАЯ СТАТИСТИКА
        ws['A5'] = "ОБЩАЯ СТАТИСТИКА"
        ws['A5'].font = title_font
        
        # Подсчет статистики
        total_items = len(items_metadata)
        camera_items = sum(1 for item in items_metadata.values() if item.get('source') == 'camera')
        video_items = total_items - camera_items
        
        # По типам предметов
        item_types = {}
        confidence_sum = 0
        timestamps = []
        
        for item in items_metadata.values():
            item_type = item.get('class_name', 'Неизвестно')
            item_types[item_type] = item_types.get(item_type, 0) + 1
            confidence_sum += item.get('confidence', 0)
            timestamp = item.get('timestamp', '')
            if timestamp:
                timestamps.append(timestamp)
        
        avg_confidence = confidence_sum / total_items if total_items > 0 else 0
        
        # Определяем период
        if timestamps:
            try:
                sorted_timestamps = sorted(timestamps)
                period = f"с {sorted_timestamps[0]} по {sorted_timestamps[-1]}"
            except:
                period = "Нет данных о времени"
        else:
            period = "Нет данных"
        
        # Таблица общей статистики
        ws['A7'] = "Показатель"
        ws['B7'] = "Значение"
        
        stats_data = [
            ["Всего обнаружено предметов", total_items],
            ["С камеры", camera_items],
            ["Из видео", video_items],
            ["Средняя уверенность", f"{avg_confidence:.1%}"],
            ["Период анализа", period]
        ]
        
        for i, (label, value) in enumerate(stats_data, start=8):
            ws[f'A{i}'] = label
            ws[f'B{i}'] = value
        
        # Форматируем заголовки
        for cell in ['A7', 'B7']:
            ws[cell].font = header_font
            ws[cell].fill = header_fill
            ws[cell].alignment = center_alignment
            ws[cell].border = border
        
        # Форматируем данные
        for row in range(8, 8 + len(stats_data)):
            for col in ['A', 'B']:
                ws[f'{col}{row}'].border = border
                ws[f'{col}{row}'].alignment = Alignment(horizontal="left" if col == 'A' else "right")
        
        # 3. РАСПРЕДЕЛЕНИЕ ПО ТИПАМ
        start_row = 8 + len(stats_data) + 2
        ws[f'A{start_row}'] = "РАСПРЕДЕЛЕНИЕ ПО ТИПАМ ПРЕДМЕТОВ"
        ws[f'A{start_row}'].font = title_font
        
        ws[f'A{start_row + 1}'] = "Тип предмета"
        ws[f'B{start_row + 1}'] = "Количество"
        ws[f'C{start_row + 1}'] = "Доля"
        
        # Заполняем данные по типам
        row = start_row + 2
        for item_type, count in sorted(item_types.items(), key=lambda x: x[1], reverse=True):
            percentage = (count / total_items * 100) if total_items > 0 else 0
            ws[f'A{row}'] = item_type
            ws[f'B{row}'] = count
            ws[f'C{row}'] = f"{percentage:.1f}%"
            row += 1
        
        if not item_types:
            ws[f'A{row}'] = "Нет данных"
            row += 1
        
        # Форматируем заголовки
        for col in ['A', 'B', 'C']:
            cell = f'{col}{start_row + 1}'
            ws[cell].font = header_font
            ws[cell].fill = header_fill
            ws[cell].alignment = center_alignment
            ws[cell].border = border
        
        # Форматируем данные
        for data_row in range(start_row + 2, row):
            for col in ['A', 'B', 'C']:
                ws[f'{col}{data_row}'].border = border
                alignment = "left" if col == 'A' else "center"
                ws[f'{col}{data_row}'].alignment = Alignment(horizontal=alignment)
        
        # 4. ПОСЛЕДНИЕ ОБНАРУЖЕННЫЕ ПРЕДМЕТЫ
        last_row = row + 2
        ws[f'A{last_row}'] = "ПОСЛЕДНИЕ ОБНАРУЖЕННЫЕ ПРЕДМЕТЫ"
        ws[f'A{last_row}'].font = title_font
        
        headers = ["Тип", "Время обнаружения", "Уверенность", "Источник", "Файл"]
        for col_idx, header in enumerate(headers, start=1):
            cell = f'{get_column_letter(col_idx)}{last_row + 1}'
            ws[cell] = header
            ws[cell].font = header_font
            ws[cell].fill = PatternFill(start_color="27ae60", end_color="27ae60", fill_type="solid")
            ws[cell].alignment = center_alignment
            ws[cell].border = border
        
        # Берем последние 10 предметов
        recent_items = list(items_metadata.items())[-10:] if items_metadata else []
        
        data_row = last_row + 2
        for filename, metadata in recent_items:
            source_text = 'Камера' if metadata.get('source') == 'camera' else 'Видео'
            confidence = metadata.get('confidence', 0)
            
            ws[f'A{data_row}'] = metadata.get('class_name', 'Неизвестно')
            ws[f'B{data_row}'] = metadata.get('timestamp', 'Нет данных')
            ws[f'C{data_row}'] = f"{confidence:.1%}"
            ws[f'D{data_row}'] = source_text
            ws[f'E{data_row}'] = filename
            
            # Цветовые метки для уверенности
            if confidence >= 0.8:
                fill_color = "d4efdf"  # зеленый
            elif confidence >= 0.6:
                fill_color = "fcf3cf"  # желтый
            else:
                fill_color = "fadbd8"  # красный
            
            ws[f'C{data_row}'].fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
            
            for col in ['A', 'B', 'C', 'D', 'E']:
                ws[f'{col}{data_row}'].border = border
                alignment = "center" if col in ['C', 'D'] else "left"
                ws[f'{col}{data_row}'].alignment = Alignment(horizontal=alignment)
            
            data_row += 1
        
        if not recent_items:
            ws[f'A{data_row}'] = "Нет данных о последних предметах"
            data_row += 1
        

        # Автонастройка ширины колонок
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # 6. ИТОГИ НА ОТДЕЛЬНОМ ЛИСТЕ
        ws2 = wb.create_sheet(title="Итоги")
        
        # Краткая сводка
        ws2['A1'] = "КРАТКАЯ СВОДКА"
        ws2['A1'].font = Font(bold=True, size=14)
        
        summary_data = [
            ["Общее количество обнаружений:", total_items],
            ["Обнаружений с камеры:", camera_items],
            ["Обнаружений из видео:", video_items],
            ["Средняя уверенность:", f"{avg_confidence:.1%}"],
            ["Уникальных типов предметов:", len(item_types)],
            ["Период анализа:", period],
            ["", ""],
            ["Самый частый тип:", max(item_types.items(), key=lambda x: x[1])[0] if item_types else "Нет данных"],
            ["Количество самого частого типа:", max(item_types.values()) if item_types else 0]
        ]
        
        for i, (label, value) in enumerate(summary_data, start=3):
            ws2[f'A{i}'] = label
            ws2[f'B{i}'] = value
            ws2[f'A{i}'].font = Font(bold=True)
        
        # График распределения (текстовый)
        if item_types:
            chart_row = len(summary_data) + 5
            ws2[f'A{chart_row}'] = "РАСПРЕДЕЛЕНИЕ ПО ТИПАМ (в штуках):"
            ws2[f'A{chart_row}'].font = Font(bold=True)
            
            for i, (item_type, count) in enumerate(sorted(item_types.items(), key=lambda x: x[1], reverse=True), start=1):
                percentage = (count / total_items * 100) if total_items > 0 else 0
                ws2[f'A{chart_row + i}'] = item_type
                ws2[f'B{chart_row + i}'] = count
                ws2[f'C{chart_row + i}'] = f"{percentage:.1f}%"
                # Простая текстовая визуализация
                bar = "█" * int((count / max(item_types.values())) * 20) if item_types else ""
                ws2[f'D{chart_row + i}'] = bar
        
        # Сохраняем в буфер
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        # Отправляем файл
        response = app.response_class(
            response=buffer.getvalue(),
            status=200,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response.headers['Content-Disposition'] = 'attachment; filename="lost_items_report.xlsx"'
        
        return response
        
    except Exception as e:
        print(f"Error generating Excel report: {e}")
        import traceback
        traceback.print_exc()
        
        return jsonify({
            'success': False, 
            'error': f'Ошибка генерации Excel отчета: {str(e)}'
        }), 500

if __name__ == '__main__':
    print(f"=== Запуск системы детекции потерянных вещей ===")
    print(f"Папка для видео: {UPLOAD_FOLDER}")
    print(f"Папка для обнаруженных предметов: {LOST_ITEMS_FOLDER}")
    print(f"Папка для кадров с камеры: {CAMERA_FRAMES_FOLDER}")
    print(f"Используется GPU: {ai_model.check_gpu()}")
    
    # Пути к SSL сертификатам
    ssl_dir = BASE_DIR / 'ssl_certs'
    ssl_cert = ssl_dir / 'cert.pem'
    ssl_key = ssl_dir / 'key.pem'
    
    # Получаем IP
    local_ip = get_ip_address()
    
    print(f"\n🔐 SSL НАСТРОЙКА")
    print("=" * 50)
    
    if ssl_cert.exists() and ssl_key.exists():
        print(f"✅ SSL сертификаты найдены:")
        print(f"   Сертификат: {ssl_cert}")
        print(f"   Ключ: {ssl_key}")
        print(f"\n📡 ДОСТУПНЫЕ АДРЕСА:")
        print(f"💻 Компьютер: https://localhost:5000")
        print(f"📱 Телефон:   https://{local_ip}:5000")
        print(f"\n⚠ ВАЖНО ДЛЯ ТЕЛЕФОНА:")
        print(f"1. Открой браузер на телефоне")
        print(f"2. Введи: https://{local_ip}:5000")
        print(f"3. Прими предупреждение о сертификате")
        print(f"4. Нажми 'Дополнительно' → 'Перейти на сайт'")
        print("=" * 50)
        
        # Создаем SSL контекст
        context = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
        context.load_cert_chain(certfile=str(ssl_cert), keyfile=str(ssl_key))
        
        # Запускаем сервер
        app.run(
            debug=True,
            host='0.0.0.0',  # Принимаем соединения со всех интерфейсов
            port=5000,
            ssl_context=context,
            threaded=True
        )
    else:
        print(f"❌ SSL сертификаты не найдены!")
        print(f"   Создайте папку ssl_certs с файлами cert.pem и key.pem")
        print(f"\n📋 Инструкция по созданию:")
        print(f"1. Открой CMD как администратор")
        print(f"2. cd {BASE_DIR}")
        print(f"3. mkdir ssl_certs")
        print(f"4. cd ssl_certs")
        print(f"5. openssl req -x509 -newkey rsa:2048 -keyout key.pem -out cert.pem -days 365 -nodes")
        print(f"\n💡 Пока запускаем без SSL (только локально):")
        print(f"   http://localhost:5000")
        print("=" * 50)
        
        app.run(
            debug=True,
            host='0.0.0.0',
            port=5000,
            threaded=True
        )